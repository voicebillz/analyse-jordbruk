# -*- coding: utf-8 -*-
"""Byggjer excel/Analyse_jordbruk_produksjonar.xlsx frå spec.py.

    python tools/build_excel.py
"""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
import spec  # noqa: E402
from formula import render, tokenize  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..", "excel", "Analyse_jordbruk_produksjonar.xlsx")

FONT = "Arial"
BLUE = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10, color="000000")
GREEN = Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
GREY = Font(name=FONT, size=9, color="595959", italic=True)
TITLE = Font(name=FONT, size=15, bold=True, color="1F3B2D")
SUB = Font(name=FONT, size=10, color="404040")
HEADF = Font(name=FONT, size=10, bold=True, color="FFFFFF")
SECF = Font(name=FONT, size=10, bold=True, color="1F3B2D")
YELLOW = PatternFill("solid", fgColor="FFFF99")
HEADFILL = PatternFill("solid", fgColor="2E5E45")
SECFILL = PatternFill("solid", fgColor="DCEADF")
KPIFILL = PatternFill("solid", fgColor="F3F7F4")
THIN = Side(style="thin", color="C8D3CB")
BOX = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

NUMFMT = {
    "kr": '#,##0;(#,##0);"-"',
    "kr2": '#,##0.00;(#,##0.00);"-"',
    "n": '#,##0;(#,##0);"-"',
    "n1": '#,##0.0;(#,##0.0);"-"',
    "n2": '#,##0.00;(#,##0.00);"-"',
    "pct": '0.0%;(0.0%);"-"',
    "txt": "@",
}

SHEETS = spec.COMMON + spec.PRODS
FIRST_ROW = 6  # første datarad
HEADERS = ["Post", "Kjelde / Tripletex-konto", "Verdi", "Eining", "Formel forklart (korleis outputen blir rekna)", "Excel-formel"]
WIDTHS = [50, 34, 16, 10, 78, 46]


def q(sheet):
    return sheet if re.fullmatch(r"[A-Za-z0-9_]+", sheet) else "'%s'" % sheet


# ------------------------------------------------------------ plassering av rader
cell = {}      # "key.id" -> (sheet, row)
label = {}     # "key.id" -> label
for sh in SHEETS:
    r = FIRST_ROW
    for row in sh["rows"]:
        if row["t"] in ("input", "calc"):
            full = "%s.%s" % (sh["key"], row["id"])
            assert full not in cell, "duplikat id " + full
            cell[full] = (sh["sheet"], r)
            label[full] = row["label"]
        r += 1
    sh["_last"] = r

# Oversikt: rad per produksjon
OV_FIRST = 7
ov_row = {p["key"]: OV_FIRST + i for i, p in enumerate(spec.PRODS)}


def make_ref(curkey, mode):
    def ref(name):
        full = name if "." in name else "%s.%s" % (curkey, name)
        if full not in cell:
            raise KeyError("Ukjend referanse %s (i %s)" % (full, curkey))
        sh, r = cell[full]
        same = sh == dict((s["key"], s["sheet"]) for s in SHEETS)[curkey]
        if mode == "xl":
            return ("C%d" % r) if same else "%s!C%d" % (q(sh), r)
        lab = label[full]
        return ("%s [C%d]" % (lab, r)) if same else ("%s [%s!C%d]" % (lab, sh, r))
    return ref


def make_active(mode):
    def active(key):
        r = ov_row[key]
        if mode == "xl":
            return 'IF(Oversikt!$B$%d="J",1,0)' % r
        return "Med i totalen %s [Oversikt!B%d → 1/0]" % (spec.PNAME[key], r)
    return active


def xl(expr, key):
    return render(tokenize(expr), make_ref(key, "xl"), make_active("xl"), "xl")


def txt(expr, key):
    return render(tokenize(expr), make_ref(key, "txt"), make_active("txt"), "txt")


def is_cross(expr, key):
    toks = tokenize(expr)
    return any(k == "id" and "." in v for k, v in toks) and all(
        (k != "id" or "." in v) for k, v in toks) and len([t for t in toks if t[0] == "op"]) == 0


# ------------------------------------------------------------ arbeidsbok
wb = Workbook()
wb.remove(wb.active)


def setup(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False


def title_block(ws, title, desc, ncols=6):
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = desc
    ws["A2"].font = SUB
    ws["A3"] = ("Fargekodar:  blå tal på gul botn = input (skal fyllast inn)   ·   svart = formel   ·   "
                "grøn = lenkje frå anna ark.   Alle kronebeløp er positive (inntekt og kostnad).")
    ws["A3"].font = GREY
    ws.row_dimensions[1].height = 22


def write_sheet(sh):
    ws = wb.create_sheet(sh["sheet"])
    setup(ws, WIDTHS)
    title_block(ws, sh["title"] + (("  –  driftsgrein " + sh["dg"]) if sh.get("dg") else ""), sh.get("desc", ""))
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=FIRST_ROW - 1, column=i, value=h)
        c.font = HEADF
        c.fill = HEADFILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "C%d" % FIRST_ROW
    r = FIRST_ROW
    for row in sh["rows"]:
        if row["t"] == "head":
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = SECFILL
            c = ws.cell(row=r, column=1, value=row["label"])
            c.font = SECF
            r += 1
            continue
        a = ws.cell(row=r, column=1, value=row["label"])
        a.font = BLACK
        a.alignment = TOP
        ws.cell(row=r, column=4, value=row["unit"] if row["unit"] != "kr/eining" else "kr").font = BLACK
        v = ws.cell(row=r, column=3)
        v.number_format = NUMFMT[row["fmt"]]
        if row["t"] == "input":
            v.value = row["val"]
            v.font = BLUE
            v.fill = YELLOW
            b = ws.cell(row=r, column=2, value=row.get("src", ""))
            b.font = GREY
            b.alignment = WRAP
            note = "Input. Hentast frå: %s." % row["src"] if row.get("src") else "Input."
            if row.get("note"):
                note += " " + row["note"]
            if row["kind"] == "info":
                note += " Er med i beløpet over, brukt berre til prisutrekning."
            e = ws.cell(row=r, column=5, value=note)
            e.font = GREY
            e.alignment = WRAP
        else:
            f = xl(row["expr"], sh["key"])
            v.value = "=" + f
            cross = is_cross(row["expr"], sh["key"])
            v.font = GREEN if cross else BLACK
            if row.get("kpi"):
                a.font = BOLD
                v.font = Font(name=FONT, size=10, bold=True, color="008000" if cross else "000000")
                for col in range(1, 7):
                    ws.cell(row=r, column=col).fill = KPIFILL
            b = ws.cell(row=r, column=2, value="Lenkje frå anna ark" if cross else "Utrekna")
            b.font = GREY
            t = "%s = %s" % (row["label"], txt(row["expr"], sh["key"]))
            if row.get("note"):
                t += "   (" + row["note"] + ")"
            e = ws.cell(row=r, column=5, value=t)
            e.font = BLACK
            e.alignment = WRAP
            fx = ws.cell(row=r, column=6, value="=" + f)
            fx.data_type = "s"
            fx.font = GREY
            fx.alignment = WRAP
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = BOX
        r += 1
    ws.print_title_rows = "%d:%d" % (FIRST_ROW - 1, FIRST_ROW - 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ------------------------------------------------------------ Les meg
ws = wb.create_sheet("Les meg")
setup(ws, [4, 34, 110])
ws["B1"] = "Nøkkeltalsanalyse jordbruk – produksjonsvis oppsett"
ws["B1"].font = TITLE
lines = [
    ("Føremål", "Arbeidsboka deler jordbruksanalysen opp i éin fane per produksjon (korn, silofôr, høy, potet, agurk, tomat, "
                "frilandsgrønsaker, gulrot, kålrot, mjølk, oksar, mellomkalv, ammekyr, smågris, purkering, slaktegris, sau, geit, egg og "
                "slaktekylling) pluss fellesark for grovfôr, faste kostnader og totaløkonomi. Oppsettet byggjer på den vedlagde "
                "analysen «Analyse 2025» (Duett-import) og kontoplanen frå Tripletex."),
    ("Oppbygging av kvar fane", "1. INPUT – produksjonsdata  →  2. INPUT – rekneskapstal (kr)  →  3. fordelte faste kostnader (lenkje)  →  "
                                "4. UTREKNING totalt  →  5. OUTPUT per eining  →  6. OUTPUT nøkkeltal. Kolonne E forklarar kvar formel i ord "
                                "med cellereferansar, kolonne F viser sjølve Excel-formelen."),
    ("Fargekodar", "Blå tal på gul botn = input som skal fyllast inn. Svart = formel. Grøn = lenkje frå eit anna ark. Nøkkeltal er utheva."),
    ("Forteikn", "Alle inntekter og kostnader blir lagt inn som positive kronebeløp. Beholdningsendring buskap/ferdigvarer: + = auke i lager "
                 "(inntekt). «Endring eigenprodusert grovfôr» følgjer kostnadsforteikn: + = nedgang i lager (kostnad)."),
    ("Med i totalen", "På fana «Oversikt» set du J/N for kvar produksjon. Berre produksjonar med J blir tekne med i Grovfôr og Totaløkonomi. "
                      "Eksempeltala for silofôr, ammekyr og sau, grovfôrareal og faste kostnader er henta frå vedlagt analyse 2025, så "
                      "nøkkeltala kan kontrollerast mot den (t.d. DB pr vf. sau u/tilskot = 315 kr). Dei andre produksjonane har "
                      "fiktive, realistiske eksempeltal og står som N."),
    ("Analysegrupper", "Eksempelanalysen har eigne ark for Mjølk <80 % / 80–120 % / >120 % påsett, Ammekyr <50 % / >50 % og Smågris / "
                       "Gris kombinert. Her ligg dette som nøkkeltalet «Analysegruppe» på kvar fane, rekna ut frå påsett- eller "
                       "framfôringsprosenten."),
    ("Grovfôr", "Fana «Grovfôr» reknar korrigert grovfôrareal (overflatedyrka/1,25 + fulldyrka + kulturbeite/2,5), «kueiningar», "
                "variabel grovfôrkostnad pr dekar og maskinkostnad pr dekar. DB pr dekar grovfôr på kvar husdyrfane fordeler arealet "
                "etter produksjonen sin del av eigen grovfôrkostnad – same metode som eksempelanalysen."),
    ("Kopling til Tripletex", "Fana «Kontokopling» viser kva Tripletex-kontoar (jordbrukskontoplanen i Tripletex) som blir summert til kvar post, "
                              "og kva produksjon ein konto høyrer til. Same kopling blir brukt i HTML-appen (app/analyse.html)."),
    ("Kueiningar og vekter", "Vektene for «kueiningar» følgjer eksempelanalysen: mjølk/oksar: kyr = 1, ungdyr over 12 md = 1/4, under 12 md = 1/8; "
                             "ammekyr: ungdyr over 12 md = 1/3, under 12 md = 1/4; sau: sum dyretal / 6; geit: sum vaksne geiter / 7. "
                             "Alle snitt av 1.1. og 31.12."),
    ("Ikkje med", "Marginalanalysane i eksempelfila (Mjølk marginal, Okse marginal, Ammekyr marginal, Marginal sau/geit) krev "
                  "tilskotssatsar og er ikkje bygde om her."),
]
r = 3
for k, v in lines:
    ws.cell(row=r, column=2, value=k).font = BOLD
    c = ws.cell(row=r, column=3, value=v)
    c.font = BLACK
    c.alignment = WRAP
    ws.row_dimensions[r].height = max(15, 15 * (len(v) // 120 + 1))
    r += 2
ws.cell(row=r, column=2, value="Faner").font = SECF
r += 1
for sh in SHEETS:
    ws.cell(row=r, column=2, value=sh["sheet"]).font = BLACK
    ws.cell(row=r, column=2).hyperlink = "#'%s'!A1" % sh["sheet"]
    ws.cell(row=r, column=2).font = Font(name=FONT, size=10, color="0563C1", underline="single")
    ws.cell(row=r, column=3, value=sh.get("desc", "")).font = BLACK
    r += 1

# ------------------------------------------------------------ Oversikt
ws = wb.create_sheet("Oversikt")
setup(ws, [24, 12, 12, 13, 15, 15, 15, 15, 15, 15, 17, 30])
title_block(ws, "Oversikt over alle produksjonar", "Summert frå kvar produksjonsfane. Set J i kolonne B for produksjonar som skal vere med i grovfôr- og totaløkonomien.")
heads = ["Produksjon", "Med i totalen (J/N)", "Eining", "Tal einingar", "Produksjons-\ninntekter", "Variable\nkostnader",
         "DB u/tilskot", "Tilskot", "DB m/tilskot", "DB u/tilskot\npr eining", "DB m/tilskot\npr eining", "DB m/tilskot etter\nmaskin og innandørsmek."]
for i, h in enumerate(heads, 1):
    c = ws.cell(row=OV_FIRST - 1, column=i, value=h)
    c.font = HEADF
    c.fill = HEADFILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[OV_FIRST - 1].height = 42
dv = DataValidation(type="list", formula1='"J,N"', allow_blank=False)
ws.add_data_validation(dv)
for p in spec.PRODS:
    rr = ov_row[p["key"]]
    sheet = q(p["sheet"])

    def C_(id):
        return "%s!C%d" % (sheet, cell["%s.%s" % (p["key"], id)][1])
    a = ws.cell(row=rr, column=1, value=p["sheet"])
    a.hyperlink = "#'%s'!A1" % p["sheet"]
    a.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    b = ws.cell(row=rr, column=2, value="J" if p["key"] in spec.EXAMPLE_ACTIVE else "N")
    b.font = BLUE
    b.fill = YELLOW
    b.alignment = Alignment(horizontal="center")
    dv.add(b)
    ws.cell(row=rr, column=3, value=p["unit_label"]).font = BLACK
    vals = [("=" + C_("n"), "n"), ("=" + C_("inn_total"), "kr"), ("=" + C_("kost_total"), "kr"), ("=" + C_("db_u"), "kr"),
            ("=" + C_("tilskot"), "kr"), ("=" + C_("db_m"), "kr"),
            ("=IF(%s=0,0,%s/%s)" % (C_("n"), C_("db_u"), C_("n")), "kr2"),
            ("=IF(%s=0,0,%s/%s)" % (C_("n"), C_("db_m"), C_("n")), "kr2"),
            ("=" + C_("db_etter"), "kr")]
    for j, (f, fmt) in enumerate(vals, 4):
        c = ws.cell(row=rr, column=j, value=f)
        c.font = GREEN
        c.number_format = NUMFMT[fmt]
    for col in range(1, 13):
        ws.cell(row=rr, column=col).border = BOX
last = OV_FIRST + len(spec.PRODS) - 1
tr = last + 1
ws.cell(row=tr, column=1, value="Sum produksjonar med J").font = BOLD
for col in (5, 6, 7, 8, 9, 12):
    L = get_column_letter(col)
    c = ws.cell(row=tr, column=col, value='=SUMIF($B$%d:$B$%d,"J",%s%d:%s%d)' % (OV_FIRST, last, L, OV_FIRST, L, last))
    c.font = BOLD
    c.number_format = NUMFMT["kr"]
    c.fill = KPIFILL
ws.freeze_panes = "B%d" % OV_FIRST
ws.cell(row=tr + 2, column=1, value="Kontroll: Sum DB m/tilskot skal vere lik «Sum DB med tilskot» på Totaløkonomi:").font = GREY
k = ws.cell(row=tr + 3, column=1, value="Avvik")
k.font = GREY
kc = ws.cell(row=tr + 3, column=9, value="=I%d-%s!C%d" % (tr, q("Totaløkonomi"), cell["total.sum_db_m"][1]))
kc.number_format = NUMFMT["kr"]
kc.font = BLACK

# ------------------------------------------------------------ fellesark + produksjonar
for sh in spec.COMMON:
    write_sheet(sh)
for p in spec.PRODS:
    write_sheet(p)

# ------------------------------------------------------------ Kontokopling
ws = wb.create_sheet("Kontokopling")
setup(ws, [22, 46, 46, 14, 40])
title_block(ws, "Kopling Tripletex-konto → post og produksjon",
            "Brukt av HTML-appen når saldobalansen blir importert. Kontoane er 4-sifra basekonto i jordbruksavdelinga (.20).")
r = 5
for i, h in enumerate(["Produksjon / ark", "Post", "Tripletex-kontoar", "Forteikn", "Merknad"], 1):
    c = ws.cell(row=r, column=i, value=h)
    c.font = HEADF
    c.fill = HEADFILL
r += 1
for sh in SHEETS:
    for row in sh["rows"]:
        if row["t"] == "input" and row["kind"] in ("tilskot", "inn", "kost", "gf", "fast") and row.get("acc"):
            ws.cell(row=r, column=1, value=sh["sheet"]).font = BLACK
            ws.cell(row=r, column=2, value=row["label"]).font = BLACK
            ws.cell(row=r, column=3, value=row["acc"]).font = BLUE
            ws.cell(row=r, column=4, value="– saldo" if row["sign"] == -1 else "+ saldo").font = BLACK
            ws.cell(row=r, column=5, value=row.get("src", "")).font = GREY
            for col in range(1, 6):
                ws.cell(row=r, column=col).border = BOX
            r += 1
r += 1
ws.cell(row=r, column=1, value="Kva produksjon ein konto høyrer til (første regel som passar vinn; første aktive produksjon i lista blir brukt)").font = SECF
r += 1
for i, h in enumerate(["Type", "Kontoar / produktkode", "Produksjon(ar)"], 1):
    c = ws.cell(row=r, column=i, value=h)
    c.font = HEADF
    c.fill = HEADFILL
r += 1
NAMES = dict(spec.PNAME, POOL="Fellespott (fordelast med nøkkel)", POOL_AREAL="Fellespott areal (nøkkel: areal / arealbehov)",
             POOL_GF="Fellespott hausta grovfôr (nøkkel: vinterfôr)", POOL_HUSDYR="Fellespott husdyr (nøkkel: direkte førte husdyrkostnader)",
             POOL_TILSK="Fellespott tilskot (nøkkel: direkte førte tilskot)", FASTE="Faste kostnader", INNMEK="Innandørsmekanisering (nøkkel)",
             TOTAL="Totaløkonomi (ikkje fordelt)")
for suf, prods in spec.SUFFIX_PROD:
    ws.cell(row=r, column=1, value="Produktkode").font = BLACK
    ws.cell(row=r, column=2, value="xxxx.20.%s" % suf).font = BLUE
    ws.cell(row=r, column=3, value=" / ".join(NAMES[x] for x in prods)).font = BLACK
    r += 1
for acc, prods in spec.ACCOUNT_PROD:
    ws.cell(row=r, column=1, value="Konto").font = BLACK
    c = ws.cell(row=r, column=2, value=acc)
    c.font = BLUE
    c.alignment = WRAP
    ws.cell(row=r, column=3, value=" / ".join(NAMES[x] for x in prods)).font = BLACK
    ws.row_dimensions[r].height = max(15, 13 * (len(acc) // 44 + 1))
    r += 1

# rekkjefølgje på faner
order = ["Les meg", "Oversikt", "Totaløkonomi", "Faste kostnader", "Grovfôr"] + [p["sheet"] for p in spec.PRODS] + ["Kontokopling"]
wb._sheets = [wb[n] for n in order]
wb.active = 1
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
# kopi til nettstaden, slik at arbeidsboka kan lastast ned derifrå
import shutil
os.makedirs(os.path.join(os.path.dirname(OUT), "..", "public"), exist_ok=True)
shutil.copy(OUT, os.path.join(os.path.dirname(OUT), "..", "public", os.path.basename(OUT)))
print("Lagra", os.path.abspath(OUT))
